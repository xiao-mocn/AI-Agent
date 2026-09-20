from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as Shape
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("0003-week-02-multiplication-foundations.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

INK = "1F3440"
TEAL = "126A70"
GOLD = "E6AA45"
PAPER = "FBF8F1"
MINT = "E8F1E8"
ROSE = "F8E9DE"
PALE = "EFF3F1"
WHITE = "FFFFFF"
MUTED = "60727B"


def rgb(value):
    return RGBColor.from_string(value)


def box(slide, kind, x, y, w, h, fill=WHITE, line=None, rounded=False):
    shape = slide.shapes.add_shape(
        Shape.ROUNDED_RECTANGLE if rounded else kind,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    return shape


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Pt(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def chip(slide, label, x, y, w, fill=GOLD, color=INK):
    box(slide, Shape.ROUNDED_RECTANGLE, x, y, w, 0.4, fill, fill, True)
    text(slide, label, x + 0.08, y + 0.09, w - 0.16, 0.2, 10, color, True, PP_ALIGN.CENTER)


def card(slide, heading, body, x, y, w, h, fill=WHITE, accent=TEAL, body_size=16):
    box(slide, Shape.ROUNDED_RECTANGLE, x, y, w, h, fill, fill, True)
    box(slide, Shape.RECTANGLE, x, y, 0.07, h, accent, accent)
    text(slide, heading, x + 0.27, y + 0.2, w - 0.48, 0.34, 16, accent, True)
    text(slide, body, x + 0.27, y + 0.68, w - 0.5, h - 0.88, body_size, INK)


def base(section):
    slide = prs.slides.add_slide(BLANK)
    box(slide, Shape.RECTANGLE, 0, 0, 13.333, 7.5, PAPER)
    text(slide, "第 2 周｜乘法里的百、十和一", 0.55, 0.28, 6.5, 0.25, 9, TEAL, True)
    text(slide, section, 8.3, 0.28, 4.45, 0.25, 9, MUTED, False, PP_ALIGN.RIGHT)
    box(slide, Shape.RECTANGLE, 0.55, 0.72, 12.2, 0.025, TEAL, TEAL)
    return slide


def marker(slide, labels, x, y, size=0.65):
    for index, label in enumerate(labels):
        left = x + index * (size + 0.35)
        box(slide, Shape.OVAL, left, y, size, size, MINT, TEAL)
        text(slide, label, left, y + 0.16, size, 0.22, 14, TEAL, True, PP_ALIGN.CENTER)


# 1. cover
s = prs.slides.add_slide(BLANK)
box(s, Shape.RECTANGLE, 0, 0, 13.333, 7.5, PAPER)
chip(s, "第二单元起步", 0.85, 0.8, 1.55)
text(s, "乘法里的百、十和一", 0.85, 1.72, 10.4, 0.7, 34, INK, True)
text(s, "两、三位数乘一位数", 0.87, 2.7, 7.1, 0.4, 19, MUTED)
text(s, "今天的任务：\n把一个数分开看，再把答案合起来。", 0.87, 4.55, 6.5, 0.8, 20, TEAL, True)
for x, label, fill in ((9.2, "100", MINT), (10.55, "20", ROSE), (11.9, "4", GOLD)):
    box(s, Shape.ROUNDED_RECTANGLE, x, 2.25, 1.05, 1.05, fill, fill, True)
    text(s, label, x, 2.59, 1.05, 0.25, 20, INK, True, PP_ALIGN.CENTER)
text(s, "× 3", 10.1, 4.15, 2.6, 0.4, 22, TEAL, True, PP_ALIGN.CENTER)

# 2. route
s = base("本周学习路线")
text(s, "两节课，完成一个乘法学习闭环", 0.85, 1.15, 11.3, 0.5, 27, INK, True)
card(s, "先复现", "混合运算和经过时间\n只做短诊断", 0.85, 2.35, 2.65, 1.65, MINT)
card(s, "再分解", "看懂百、十和一\n分别表示什么", 3.85, 2.35, 2.65, 1.65, WHITE, GOLD)
card(s, "再检测", "独立完成、解释理由\n并估一估", 6.85, 2.35, 2.65, 1.65, ROSE, GOLD)
card(s, "最后预热", "先估再算\n不提前学新算法", 9.85, 2.35, 2.65, 1.65, MINT)
text(s, "只有解释得清楚，才算真的会。", 0.88, 5.0, 11.3, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 3. review
s = base("第 1 节｜已学内容复现")
chip(s, "请独立完成", 0.78, 1.0, 1.35)
text(s, "先写答案，再写第一步做了什么", 0.78, 1.63, 10.8, 0.45, 25, INK, True)
card(s, "题 1｜运算顺序", "18 − 3 × 4 = ____\n你先算哪一步？", 1.05, 2.7, 4.7, 1.65, MINT)
card(s, "题 2｜经过时间", "图书馆 8:30 开馆，\n17:00 闭馆。开放多久？", 6.9, 2.7, 4.7, 1.65, ROSE, GOLD)
text(s, "不要急着说答案，把你的想法留在纸上。", 1.05, 5.2, 10.7, 0.35, 17, TEAL, True)

# 4. time diagnostic
s = base("一天的时间｜快速诊断")
text(s, "在时间线上标出开馆和闭馆时刻", 0.88, 1.15, 11.4, 0.45, 25, INK, True, PP_ALIGN.CENTER)
box(s, Shape.RECTANGLE, 1.3, 3.7, 10.7, 0.01, TEAL, TEAL)
for hour in range(8, 18):
    x = 1.3 + (hour - 8) * 1.07
    box(s, Shape.RECTANGLE, x, 3.52, 0.01, 0.36, TEAL, TEAL)
    text(s, str(hour), x - 0.16, 4.0, 0.34, 0.25, 12, MUTED, False, PP_ALIGN.CENTER)
text(s, "8:30", 1.64, 2.75, 0.75, 0.3, 18, TEAL, True, PP_ALIGN.CENTER)
text(s, "17:00", 10.62, 2.75, 0.85, 0.3, 18, GOLD, True, PP_ALIGN.CENTER)
text(s, "从 8:30 到 17:00，经过多长时间？", 1.2, 5.15, 10.9, 0.35, 19, INK, True, PP_ALIGN.CENTER)

# 5. diagnostic feedback
s = base("第 1 节｜复现反馈")
text(s, "先说清楚“从哪里到哪里”", 0.88, 1.15, 11.4, 0.45, 26, INK, True, PP_ALIGN.CENTER)
card(s, "运算顺序", "18 − 3 × 4\n= 18 − 12\n= 6", 1.2, 2.35, 3.2, 2.05, MINT, body_size=20)
card(s, "经过时间", "8:30 → 17:00\n8 小时 30 分", 5.05, 2.35, 3.2, 2.05, WHITE, GOLD, 20)
card(s, "检查", "结果要小于\n从 8:00 到 17:00 的 9 小时", 8.9, 2.35, 3.2, 2.05, ROSE, GOLD, 16)
text(s, "已学内容只用来复现，不重新讲一遍。", 1.2, 5.25, 10.9, 0.32, 17, TEAL, True, PP_ALIGN.CENTER)

# 6. olympiad warm-up
s = base("数数图形｜观察热身")
text(s, "5 个点在一条线上，能组成多少条线段？", 0.85, 1.2, 11.4, 0.45, 25, INK, True, PP_ALIGN.CENTER)
box(s, Shape.RECTANGLE, 2.15, 3.3, 8.9, 0.01, TEAL, TEAL)
marker(s, ["A", "B", "C", "D", "E"], 2.0, 3.0, 0.65)
text(s, "要求：按顺序记录，怎样才能不重复、不遗漏？", 1.1, 5.15, 11.0, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 7. warm-up feedback
s = base("数数图形｜热身反馈")
text(s, "从左边的点开始，依次记录", 0.9, 1.15, 11.3, 0.45, 25, INK, True, PP_ALIGN.CENTER)
card(s, "从 A 出发", "AB、AC、AD、AE\n4 条", 1.1, 2.45, 2.45, 1.7, MINT)
card(s, "从 B 出发", "BC、BD、BE\n3 条", 3.95, 2.45, 2.45, 1.7, WHITE, GOLD)
card(s, "从 C、D 出发", "CD、CE、DE\n2 + 1 条", 6.8, 2.45, 2.45, 1.7, ROSE, GOLD)
card(s, "一共", "4 + 3 + 2 + 1\n= 10 条", 9.65, 2.45, 2.45, 1.7, MINT)
text(s, "今天只借用这个方法：有顺序地记录，才能看清结构。", 1.1, 5.2, 11.0, 0.32, 17, TEAL, True, PP_ALIGN.CENTER)

# 8. main example
s = base("第 1 节｜新内容主例题")
text(s, "学校买 3 盒彩笔，每盒 124 支，一共多少支？", 0.82, 1.18, 11.7, 0.5, 25, INK, True, PP_ALIGN.CENTER)
for index in range(3):
    x = 2.1 + index * 3.15
    box(s, Shape.ROUNDED_RECTANGLE, x, 2.55, 2.35, 1.75, MINT, MINT, True)
    text(s, "第 " + str(index + 1) + " 盒", x, 2.82, 2.35, 0.3, 15, TEAL, True, PP_ALIGN.CENTER)
    text(s, "124 支", x, 3.43, 2.35, 0.35, 23, INK, True, PP_ALIGN.CENTER)
text(s, "问题：3 个 124 合起来是多少？", 1.2, 5.18, 10.9, 0.35, 19, TEAL, True, PP_ALIGN.CENTER)

# 9. silent thinking
s = base("静想一下")
text(s, "先不列竖式，先把 124 看清楚", 0.85, 1.28, 11.6, 0.5, 28, INK, True, PP_ALIGN.CENTER)
card(s, "想一想", "124 里面有几个百？\n几个十？几个一？", 1.35, 2.7, 4.7, 1.8, MINT, body_size=20)
card(s, "画一画", "如果每一部分都乘 3，\n该怎样写？", 7.25, 2.7, 4.7, 1.8, ROSE, GOLD, 20)
text(s, "把你的分法写在草稿纸上。", 1.0, 5.4, 11.2, 0.34, 18, TEAL, True, PP_ALIGN.CENTER)

# 10. decomposition
s = base("第 1 节｜分解来看")
text(s, "124 = 100 + 20 + 4", 0.85, 1.1, 11.6, 0.5, 28, INK, True, PP_ALIGN.CENTER)
for x, top, bottom, fill in ((1.25, "100 × 3", "300", MINT), (4.85, "20 × 3", "60", ROSE), (8.45, "4 × 3", "12", GOLD)):
    box(s, Shape.ROUNDED_RECTANGLE, x, 2.35, 2.65, 2.05, fill, fill, True)
    text(s, top, x, 2.74, 2.65, 0.3, 20, INK, True, PP_ALIGN.CENTER)
    text(s, bottom, x, 3.45, 2.65, 0.38, 26, TEAL, True, PP_ALIGN.CENTER)
text(s, "3 盒里有 3 个百、6 个十、12 个一。", 1.0, 5.25, 11.3, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 11. combine
s = base("第 1 节｜合起来算")
text(s, "300 + 60 + 12 = 372", 0.85, 1.15, 11.6, 0.5, 29, INK, True, PP_ALIGN.CENTER)
card(s, "百位", "100 × 3 = 300\n3 个百", 1.15, 2.55, 3.15, 1.65, MINT, body_size=18)
card(s, "十位", "20 × 3 = 60\n6 个十", 5.1, 2.55, 3.15, 1.65, WHITE, GOLD, 18)
card(s, "个位", "4 × 3 = 12\n12 个一", 9.05, 2.55, 3.15, 1.65, ROSE, GOLD, 18)
text(s, "答：一共 372 支彩笔。", 1.0, 5.25, 11.3, 0.35, 19, TEAL, True, PP_ALIGN.CENTER)

# 12. estimate
s = base("第 1 节｜估一估，再检查")
text(s, "先把 124 看成接近的 120", 0.85, 1.15, 11.6, 0.45, 26, INK, True, PP_ALIGN.CENTER)
box(s, Shape.ROUNDED_RECTANGLE, 1.6, 2.55, 3.1, 1.55, MINT, MINT, True)
text(s, "120 × 3 = 360", 1.6, 3.02, 3.1, 0.35, 22, TEAL, True, PP_ALIGN.CENTER)
text(s, "→", 5.1, 2.92, 0.9, 0.4, 26, GOLD, True, PP_ALIGN.CENTER)
box(s, Shape.ROUNDED_RECTANGLE, 6.35, 2.55, 3.1, 1.55, ROSE, ROSE, True)
text(s, "准确结果 372", 6.35, 3.02, 3.1, 0.35, 22, INK, True, PP_ALIGN.CENTER)
text(s, "372 比 360 多一点，很合理。", 1.1, 5.23, 11.1, 0.35, 19, TEAL, True, PP_ALIGN.CENTER)

# 13. method card
s = base("本周方法卡")
text(s, "把数分成几个百、几个十和几个一", 0.78, 1.12, 11.8, 0.5, 26, INK, True, PP_ALIGN.CENTER)
card(s, "第 1 步", "分开看\n124 = 100 + 20 + 4", 1.0, 2.45, 3.2, 1.8, MINT, body_size=17)
card(s, "第 2 步", "分别乘\n100×3、20×3、4×3", 5.05, 2.45, 3.2, 1.8, WHITE, GOLD, 17)
card(s, "第 3 步", "合起来，再估一估\n300 + 60 + 12", 9.1, 2.45, 3.2, 1.8, ROSE, GOLD, 17)
text(s, "先估一估，能帮我们发现明显算错的结果。", 1.0, 5.25, 11.3, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 14. independent check
s = base("第 2 节｜独立检测")
chip(s, "独立完成", 0.8, 1.0, 1.25)
text(s, "213 × 4 = ______", 0.8, 1.65, 11.7, 0.55, 30, INK, True, PP_ALIGN.CENTER)
card(s, "请写两件事", "① 先估一估结果大约是多少。\n② 任选“分解”或“竖式”说明你的计算。", 2.05, 3.0, 9.25, 1.65, MINT, body_size=18)
text(s, "完成后检查：结果应不应大约在 800 左右？", 1.1, 5.35, 11.1, 0.32, 18, TEAL, True, PP_ALIGN.CENTER)

# 15. check feedback
s = base("第 2 节｜检测反馈")
text(s, "213 × 4 = 852", 0.85, 1.12, 11.6, 0.5, 29, INK, True, PP_ALIGN.CENTER)
text(s, "200 × 4 + 10 × 4 + 3 × 4", 1.15, 2.35, 11.0, 0.35, 19, MUTED, True, PP_ALIGN.CENTER)
card(s, "800", "4 个百", 1.3, 3.1, 2.45, 1.35, MINT, body_size=20)
card(s, "40", "4 个十", 5.45, 3.1, 2.45, 1.35, WHITE, GOLD, 20)
card(s, "12", "12 个一", 9.6, 3.1, 2.45, 1.35, ROSE, GOLD, 20)
text(s, "800 + 40 + 12 = 852；和 200 × 4 约等于 800 相符。", 1.05, 5.25, 11.2, 0.35, 17, TEAL, True, PP_ALIGN.CENTER)

# 16. transfer question
s = base("第 2 节｜迁移检测")
text(s, "礼堂每排有 132 个座位，共 3 排。", 0.9, 1.25, 11.5, 0.4, 25, INK, True, PP_ALIGN.CENTER)
text(s, "一共有多少个座位？", 0.9, 1.9, 11.5, 0.4, 25, INK, True, PP_ALIGN.CENTER)
card(s, "先估", "把 132 看成 130，\n130 × 3 大约是多少？", 1.2, 3.05, 4.6, 1.65, MINT, body_size=18)
card(s, "再解释", "132 里的 2 × 3，\n在题目里表示什么？", 7.0, 3.05, 4.6, 1.65, ROSE, GOLD, 18)

# 17. transfer feedback
s = base("第 2 节｜迁移反馈")
text(s, "132 × 3 = 396", 0.85, 1.13, 11.6, 0.5, 29, INK, True, PP_ALIGN.CENTER)
card(s, "分解", "100×3 + 30×3 + 2×3\n= 300 + 90 + 6", 1.15, 2.5, 4.6, 1.85, MINT, body_size=18)
card(s, "解释个位", "2 × 3 = 6\n每排个位上的 2 个座位\n重复 3 次", 7.0, 2.5, 4.6, 1.85, ROSE, GOLD, 16)
text(s, "答：一共有 396 个座位。先估 130 × 3 约是 390，也合理。", 0.95, 5.25, 11.4, 0.35, 17, TEAL, True, PP_ALIGN.CENTER)

# 18. preview
s = base("第 2 节｜课末预习")
chip(s, "只预测，不讲新算法", 0.75, 1.0, 1.9)
text(s, "每箱 206 瓶水，4 箱大约有多少瓶？", 0.85, 1.7, 11.6, 0.45, 26, INK, True, PP_ALIGN.CENTER)
for index in range(4):
    x = 1.65 + index * 2.65
    box(s, Shape.ROUNDED_RECTANGLE, x, 3.0, 1.9, 1.35, MINT if index % 2 == 0 else ROSE,
        MINT if index % 2 == 0 else ROSE, True)
    text(s, "206 瓶", x, 3.48, 1.9, 0.28, 18, INK, True, PP_ALIGN.CENTER)
text(s, "预测问题：你会把 206 看成多少来估？先说理由，不用列竖式。", 0.9, 5.25, 11.5, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

prs.core_properties.title = "第 2 周｜乘法里的百、十和一"
prs.core_properties.subject = "三年级数学学生课堂演示稿"
prs.core_properties.author = "Codex"
prs.save(OUT)
print(f"created {OUT} with {len(prs.slides)} slides")
