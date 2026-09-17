"""Guard against passing a color string where a layout helper expects a height."""
import ast
from pathlib import Path

SOURCE = Path(__file__).with_name("build-week01-ppt.py")
tree = ast.parse(SOURCE.read_text(encoding="utf-8"))
invalid = []

for node in ast.walk(tree):
    if not (isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "card"):
        continue
    if len(node.args) < 7 or not isinstance(node.args[6], ast.Constant) or not isinstance(node.args[6].value, (int, float)):
        invalid.append(node.lineno)

assert not invalid, f"card() 的第 7 个位置参数必须是数值高度，异常行：{invalid}"

from pptx import Presentation

DECK = Path(__file__).with_name("0002-week-01-quantity-modeling-and-patterns.pptx")
if DECK.exists():
    presentation = Presentation(DECK)
    count = len(presentation.slides)
    assert 14 <= count <= 20, f"紧凑双课时 PPT 页数应为 14–20 页，当前为 {count} 页"
    deck_text = "\n".join(
        shape.text for slide in presentation.slides for shape in slide.shapes
        if hasattr(shape, "text")
    )
    assert deck_text.count("经典例题：") == 1, "每周合并 PPT 只能有 1 道完整经典例题"
    for marker in (
        "第 1 节｜易错点检测", "第 1 节｜新内容经典例题", "静想一下", "第 1 节｜例题分析",
        "第 2 节｜新内容检测", "第 2 节｜检测反馈", "第 2 节｜迁移检测", "第 2 节｜课末预习",
    ):
        assert marker in deck_text, f"缺少连续双课时环节：{marker}"
    for forbidden in ("教师版", "教师复盘", "记录学生", "下次调整", "教师提示"):
        assert forbidden not in deck_text, f"学生 PPT 不应包含教师内容：{forbidden}"
    slide_texts = ["\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")) for slide in presentation.slides]
    def index_of(marker):
        return next(index for index, value in enumerate(slide_texts) if marker in value)
    assert index_of("第 1 节｜易错点检测") < index_of("第 1 节｜新内容经典例题") < index_of("第 2 节｜新内容检测")
    assert index_of("第 2 节｜新内容检测") < index_of("第 2 节｜检测反馈") < index_of("第 2 节｜迁移检测") < index_of("第 2 节｜课末预习")
