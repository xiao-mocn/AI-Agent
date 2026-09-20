from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as Shape
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("0004-week-02-addition-subtraction-strategies.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

INK, TEAL, GOLD = "1F3440", "126A70", "E6AA45"
PAPER, MINT, ROSE, WHITE, MUTED = "FBF8F1", "E8F1E8", "F8E9DE", "FFFFFF", "60727B"


def rgb(value):
    return RGBColor.from_string(value)


def shape(slide, kind, x, y, w, h, fill=WHITE, line=None, rounded=False):
    item = slide.shapes.add_shape(
        Shape.ROUNDED_RECTANGLE if rounded else kind,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    item.fill.solid()
    item.fill.fore_color.rgb = rgb(fill)
    item.line.color.rgb = rgb(line or fill)
    return item


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = item.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Pt(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return item


def base(section):
    slide = prs.slides.add_slide(BLANK)
    shape(slide, Shape.RECTANGLE, 0, 0, 13.333, 7.5, PAPER)
    text(slide, "第 2 周｜看见结构，巧算加减", 0.55, 0.28, 6.5, 0.25, 9, TEAL, True)
    text(slide, section, 8.2, 0.28, 4.55, 0.25, 9, MUTED, False, PP_ALIGN.RIGHT)
    shape(slide, Shape.RECTANGLE, 0.55, 0.72, 12.2, 0.025, TEAL, TEAL)
    return slide


def card(slide, heading, body, x, y, w, h, fill=WHITE, accent=TEAL, body_size=16):
    shape(slide, Shape.ROUNDED_RECTANGLE, x, y, w, h, fill, fill, True)
    shape(slide, Shape.RECTANGLE, x, y, 0.07, h, accent, accent)
    text(slide, heading, x + 0.27, y + 0.2, w - 0.48, 0.34, 16, accent, True)
    text(slide, body, x + 0.27, y + 0.68, w - 0.5, h - 0.88, body_size, INK)


def chip(slide, label, x, y, w):
    shape(slide, Shape.ROUNDED_RECTANGLE, x, y, w, 0.4, GOLD, GOLD, True)
    text(slide, label, x + 0.08, y + 0.09, w - 0.16, 0.2, 10, INK, True, PP_ALIGN.CENTER)


# 1 cover
s = prs.slides.add_slide(BLANK)
shape(s, Shape.RECTANGLE, 0, 0, 13.333, 7.5, PAPER)
chip(s, "思维专题", 0.85, 0.8, 1.2)
text(s, "看见结构，巧算加减", 0.85, 1.72, 10.6, 0.7, 34, INK, True)
text(s, "第 2 周｜加减巧算", 0.87, 2.68, 6.4, 0.4, 19, MUTED)
text(s, "今天不比谁算得快，\n比谁先看见可以凑整的结构。", 0.87, 4.55, 6.6, 0.8, 20, TEAL, True)
for x, number, fill in ((9.15, "457", MINT), (10.7, "143", ROSE), (12.05, "600", GOLD)):
    shape(s, Shape.ROUNDED_RECTANGLE, x, 2.35, 1.1, 0.88, fill, fill, True)
    text(s, number, x, 2.65, 1.1, 0.22, 17, INK, True, PP_ALIGN.CENTER)
text(s, "＋", 10.28, 2.62, 0.35, 0.2, 18, TEAL, True, PP_ALIGN.CENTER)
text(s, "＝", 11.66, 2.62, 0.3, 0.2, 18, TEAL, True, PP_ALIGN.CENTER)

# 2 route
s = base("本周学习路线")
text(s, "一节思维课：复现 → 发现 → 迁移 → 预热", 0.85, 1.15, 11.5, 0.5, 26, INK, True)
card(s, "短复习", "校内易错点\n只复现，不新授", 0.85, 2.35, 2.65, 1.65, MINT)
card(s, "主专题", "加减巧算\n看见凑整结构", 3.85, 2.35, 2.65, 1.65, WHITE, GOLD)
card(s, "独立迁移", "改写算式\n说清为什么", 6.85, 2.35, 2.65, 1.65, ROSE, GOLD)
card(s, "下一次", "巧添符号\n先猜想，不讲解", 9.85, 2.35, 2.65, 1.65, MINT)
text(s, "每一次变形，都要说清“为什么还相等”。", 0.85, 5.05, 11.5, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 3 review
s = base("第 1 节｜课内综合诊断")
chip(s, "独立完成", 0.8, 1.0, 1.2)
text(s, "只用几分钟，找出今天要注意的地方", 0.8, 1.63, 11.5, 0.45, 24, INK, True)
card(s, "混合运算", "(96 − 24) ÷ 8 + 13 × 2 = ____\n写出每一步的依据。", 1.05, 2.72, 4.7, 1.65, MINT)
card(s, "时间综合", "9:00 出发，行程共 6 小时 50 分。\n16:00 返回，还剩几分钟？", 6.9, 2.72, 4.7, 1.65, ROSE, GOLD)
text(s, "课内问题只用来诊断，不占用今天的主专题。", 1.05, 5.22, 10.8, 0.35, 17, TEAL, True)

# 4 review feedback
s = base("第 1 节｜复习反馈")
text(s, "综合复习：算对，更要能检查", 0.85, 1.2, 11.6, 0.45, 25, INK, True, PP_ALIGN.CENTER)
card(s, "混合运算", "(96 − 24) ÷ 8 + 13 × 2\n= 9 + 26 = 35", 1.45, 2.45, 3.2, 2.0, MINT, body_size=18)
card(s, "时间综合", "从 9:00 到 16:00 是 7 小时。\n7 小时 − 6 小时 50 分\n= 10 分", 5.1, 2.45, 3.2, 2.0, WHITE, GOLD, 16)
card(s, "带到后面", "巧算也不能随便换号、\n丢掉括号或改变减数。", 8.75, 2.45, 3.2, 2.0, ROSE, GOLD, 16)

# 5 figure warm-up
s = base("数数图形｜复习热身")
text(s, "7 个点在一条线上，有多少条线段？", 0.85, 1.2, 11.5, 0.45, 25, INK, True, PP_ALIGN.CENTER)
shape(s, Shape.RECTANGLE, 2.15, 3.32, 8.9, 0.02, TEAL, TEAL)
for index, label in enumerate(("A", "B", "C", "D", "E", "F", "G")):
    x = 1.7 + index * 0.85
    shape(s, Shape.OVAL, x, 3.0, 0.6, 0.6, MINT, TEAL)
    text(s, label, x, 3.17, 0.6, 0.2, 13, TEAL, True, PP_ALIGN.CENTER)
text(s, "只问一件事：怎样记录才不会重复、不会遗漏？", 1.1, 5.15, 11.1, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 6 warm-up feedback
s = base("数数图形｜热身反馈")
text(s, "从左到右，有顺序地记录", 0.9, 1.15, 11.4, 0.45, 25, INK, True, PP_ALIGN.CENTER)
card(s, "从 A、B 出发", "6 + 5 条", 1.3, 2.55, 2.3, 1.25, MINT, body_size=20)
card(s, "从 C、D 出发", "4 + 3 条", 4.05, 2.55, 2.3, 1.25, WHITE, GOLD, 20)
card(s, "从 E、F 出发", "2 + 1 条", 6.8, 2.55, 2.3, 1.25, ROSE, GOLD, 20)
card(s, "一共", "6+5+4+3+2+1\n= 21 条", 9.55, 2.55, 2.3, 1.25, MINT, body_size=16)
text(s, "有序枚举不是慢，而是用结构保证完整。", 1.0, 5.12, 11.3, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 7 main example
s = base("第 1 节｜主例题")
text(s, "457 + 321 + 143 + 479", 0.85, 1.35, 11.6, 0.55, 31, INK, True, PP_ALIGN.CENTER)
text(s, "不用从左到右马上计算。先观察：哪两个数相加最接近整百？", 0.95, 2.45, 11.3, 0.35, 19, TEAL, True, PP_ALIGN.CENTER)
card(s, "任务", "圈出你想先相加的一对数，\n并写出理由。", 2.1, 3.5, 9.1, 1.45, MINT, body_size=19)

# 8 silent think
s = base("静想一下")
text(s, "先不动笔算，先找“朋友数”", 0.85, 1.28, 11.6, 0.5, 28, INK, True, PP_ALIGN.CENTER)
card(s, "试一试", "457 和谁相加\n最接近整百？", 1.35, 2.75, 4.65, 1.7, MINT, body_size=20)
card(s, "再想一想", "321 和谁相加\n最接近整百？", 7.25, 2.75, 4.65, 1.7, ROSE, GOLD, 20)
text(s, "先写猜想，再用计算验证。", 1.0, 5.4, 11.2, 0.34, 18, TEAL, True, PP_ALIGN.CENTER)

# 9 pairing
s = base("第 1 节｜看见结构")
text(s, "457 + 143 = 600", 1.0, 1.25, 5.1, 0.45, 27, TEAL, True, PP_ALIGN.CENTER)
text(s, "321 + 479 = 800", 7.2, 1.25, 5.1, 0.45, 27, GOLD, True, PP_ALIGN.CENTER)
card(s, "可以这样排", "(457 + 143) + (321 + 479)", 1.3, 2.85, 10.75, 1.25, MINT, body_size=23)
text(s, "加法交换律和结合律：加数换位置、重新分组，和不变。", 1.0, 5.05, 11.3, 0.35, 18, INK, True, PP_ALIGN.CENTER)

# 10 solution
s = base("第 1 节｜规范解答")
text(s, "457 + 321 + 143 + 479", 1.05, 1.2, 10.9, 0.35, 24, INK, True, PP_ALIGN.CENTER)
text(s, "= (457 + 143) + (321 + 479)", 1.05, 2.15, 10.9, 0.35, 22, TEAL, True, PP_ALIGN.CENTER)
text(s, "= 600 + 800", 1.05, 3.1, 10.9, 0.35, 22, TEAL, True, PP_ALIGN.CENTER)
text(s, "= 1400", 1.05, 4.05, 10.9, 0.35, 26, GOLD, True, PP_ALIGN.CENTER)
text(s, "先找能凑整的数，再用规则把它们放在一起。", 1.0, 5.35, 11.3, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 11 plus minus
s = base("第 1 节｜减加也能凑整")
text(s, "623 − 251 + 177", 0.85, 1.2, 11.6, 0.5, 29, INK, True, PP_ALIGN.CENTER)
card(s, "先看", "623 和 177\n可以凑成多少？", 1.3, 2.75, 4.6, 1.7, MINT, body_size=21)
card(s, "再算", "(623 + 177) − 251\n= 800 − 251\n= 549", 7.1, 2.75, 4.6, 1.7, ROSE, GOLD, 19)
text(s, "能调整的是加数的位置；减去 251 的意思不能丢。", 1.0, 5.3, 11.3, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 12 continuous subtraction
s = base("第 1 节｜连续减法")
text(s, "537 − 158 − 42", 0.85, 1.2, 11.6, 0.5, 29, INK, True, PP_ALIGN.CENTER)
card(s, "先合并减数", "158 + 42 = 200", 1.25, 2.75, 4.5, 1.45, MINT, body_size=22)
card(s, "再从总数里减", "537 − (158 + 42)\n= 537 − 200\n= 337", 7.05, 2.75, 4.75, 1.45, ROSE, GOLD, 18)
text(s, "连续减去两个数，等于减去这两个数的和。", 1.0, 5.3, 11.3, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 13 method card
s = base("本周方法卡｜加减巧算")
card(s, "找结构", "谁和谁相加\n能凑整？", 1.0, 2.2, 3.2, 1.85, MINT, body_size=20)
card(s, "按规则变形", "加数可交换、结合；\n减数不能随意变号。", 5.05, 2.2, 3.2, 1.85, WHITE, GOLD, 18)
card(s, "算完检验", "估一估、反着算，\n看结果是否合理。", 9.1, 2.2, 3.2, 1.85, ROSE, GOLD, 18)
text(s, "巧算不是跳步，而是让每一步更有理由。", 1.0, 5.25, 11.3, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 14 independent check
s = base("第 2 节｜独立检测")
chip(s, "独立完成", 0.8, 1.0, 1.2)
text(s, "999 − 1 + 998 − 2 + 997 − 3 + 996 − 4", 0.65, 1.75, 12.0, 0.5, 24, INK, True, PP_ALIGN.CENTER)
card(s, "要求", "① 自己决定怎样分组。\n② 写出每一组的共同结构。\n③ 用估算检查最后结果。", 2.0, 3.05, 9.3, 1.6, MINT, body_size=19)
text(s, "不要只写答案，要把“看见的规律”写出来。", 1.0, 5.38, 11.3, 0.32, 18, TEAL, True, PP_ALIGN.CENTER)

# 15 feedback
s = base("第 2 节｜检测反馈")
text(s, "999 − 1 + 998 − 2 + 997 − 3 + 996 − 4", 0.65, 1.15, 12.0, 0.45, 22, INK, True, PP_ALIGN.CENTER)
card(s, "前两组", "999−1=998\n998−2=996", 1.25, 2.55, 4.45, 1.45, MINT, body_size=19)
card(s, "后两组", "997−3=994\n996−4=992", 7.0, 2.55, 4.45, 1.45, ROSE, GOLD, 19)
text(s, "998 + 996 + 994 + 992 = 3980", 1.0, 4.65, 11.3, 0.45, 24, TEAL, True, PP_ALIGN.CENTER)
text(s, "检查：4 个接近 1000 的数相加，结果接近 4000，合理。", 1.0, 5.35, 11.3, 0.32, 17, MUTED, True, PP_ALIGN.CENTER)

# 16 transfer question
s = base("第 2 节｜迁移检测")
text(s, "1000 − 81 − 19 − 82 − 18 − 83 − 17", 0.65, 1.25, 12.0, 0.5, 26, INK, True, PP_ALIGN.CENTER)
card(s, "先找配对", "哪些减数相加\n正好是 100？", 1.3, 2.9, 4.6, 1.65, MINT, body_size=20)
card(s, "再判断", "连续减法能否改写成\n减去一串减数的和？", 7.1, 2.9, 4.6, 1.65, ROSE, GOLD, 19)
text(s, "先写出改写后的算式，再计算。", 1.0, 5.35, 11.3, 0.32, 18, TEAL, True, PP_ALIGN.CENTER)

# 17 transfer feedback
s = base("第 2 节｜迁移反馈")
text(s, "1000 − 81 − 19 − 82 − 18 − 83 − 17", 0.65, 1.2, 12.0, 0.35, 23, INK, True, PP_ALIGN.CENTER)
text(s, "= 1000 − [(81+19) + (82+18) + (83+17)]", 0.65, 2.05, 12.0, 0.35, 19, TEAL, True, PP_ALIGN.CENTER)
text(s, "= 1000 − 300", 1.0, 2.9, 11.3, 0.35, 22, TEAL, True, PP_ALIGN.CENTER)
text(s, "= 700", 1.0, 3.75, 11.3, 0.35, 26, GOLD, True, PP_ALIGN.CENTER)
text(s, "连续减去多个数，先把所有减数相加，再从总数里减。", 1.0, 5.15, 11.3, 0.35, 18, TEAL, True, PP_ALIGN.CENTER)

# 18 preview
s = base("第 2 节｜课末预习")
chip(s, "巧添符号｜预热", 0.78, 1.0, 1.65)
text(s, "1  2  3  4  5  =  10", 0.85, 1.78, 11.6, 0.5, 30, INK, True, PP_ALIGN.CENTER)
card(s, "猜一猜", "在相邻数字之间，可以添哪些运算符号或括号？", 1.45, 3.0, 4.65, 1.55, MINT, body_size=19)
card(s, "从目标倒着想", "要得到 10，最后一步可能是什么？\n先凑接近 10 的数也可以。", 7.2, 3.0, 4.65, 1.55, ROSE, GOLD, 18)
text(s, "今天只记录猜想；下一次再学习怎样有顺序地尝试。", 1.0, 5.35, 11.3, 0.32, 18, TEAL, True, PP_ALIGN.CENTER)

prs.core_properties.title = "第 2 周｜看见结构，巧算加减"
prs.core_properties.subject = "三年级数学学生课堂演示稿"
prs.core_properties.author = "Codex"
prs.save(OUT)
print(f"created {OUT} with {len(prs.slides)} slides")
