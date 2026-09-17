from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as Shape
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

OUT = Path(r"E:\project\AI-Agent\courses\primary-math-tutoring\slides\0002-week-01-quantity-modeling-and-patterns.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

INK = "1F3440"; TEAL = "126A70"; GOLD = "E6AA45"; PAPER = "FBF8F1"; MINT = "E8F1E8"; ROSE = "F8E9DE"; PALE = "EFF3F1"; WHITE = "FFFFFF"; MUTED = "60727B"

def rgb(value): return RGBColor.from_string(value)
def shape(slide, kind, x, y, w, h, fill=PAPER, line=None, radius=False):
    shp = slide.shapes.add_shape(Shape.ROUNDED_RECTANGLE if radius else kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill); shp.line.color.rgb = rgb(line or fill)
    return shp
def text(slide, value, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = value; run.font.name = "Arial"; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(0)
    return box
def base(section=""):
    s = prs.slides.add_slide(BLANK); shape(s, Shape.RECTANGLE, 0, 0, 13.333, 7.5, PAPER)
    text(s, "第 1 周｜看见关系，发现规律", .55, .28, 6.4, .25, 9, TEAL, True)
    if section: text(s, section, 8.7, .28, 4.05, .25, 9, MUTED, False, PP_ALIGN.RIGHT)
    shape(s, Shape.RECTANGLE, .55, .7, 12.2, .025, TEAL)
    return s
def card(slide, heading, body, x, y, w, h, fill=WHITE, accent=TEAL):
    shape(slide, Shape.ROUNDED_RECTANGLE, x, y, w, h, fill, fill, True); shape(slide, Shape.RECTANGLE, x, y, .07, h, accent, accent)
    text(slide, heading, x+.25, y+.22, w-.45, .34, 16, accent, True)
    text(slide, body, x+.25, y+.72, w-.48, h-.9, 13, INK)
def chip(slide, label, x, y, w, fill, color=INK):
    shape(slide, Shape.ROUNDED_RECTANGLE, x, y, w, .37, fill, fill, True); text(slide, label, x+.08, y+.08, w-.16, .18, 10, color, True, PP_ALIGN.CENTER)

# 1. cover
s = prs.slides.add_slide(BLANK); shape(s, Shape.RECTANGLE, 0, 0, 13.333, 7.5, PAPER)
chip(s, "数学挑战", .8, .8, 1.2, GOLD)
text(s, "先复习，再学新方法", .8, 1.8, 10.5, .8, 34, INK, True)
text(s, "第 1 周｜连续两节课", .82, 2.85, 5.5, .4, 18, MUTED)
text(s, "今天只讲透一个方法：\n先把数量对齐，再找出未知量。", .82, 4.85, 6.2, .75, 20, TEAL, True)
shape(s, Shape.OVAL, 10.5, 1.7, 1.5, 1.5, MINT, MINT); shape(s, Shape.OVAL, 11.45, 3.0, .72, .72, ROSE, ROSE)

# 2. weekly flow
s = base("本周学习路线")
text(s, "两节课，完成一个学习闭环", .8, 1.15, 11.0, .5, 27, INK, True)
card(s, "第 1 节", "易错点检测\n→ 学习新方法", .95, 2.3, 5.2, 2.05, MINT)
card(s, "第 2 节", "独立检测与订正\n→ 预习下一内容", 7.15, 2.3, 5.2, 2.05, ROSE, GOLD)
text(s, "先把旧问题找出来，再把新方法学明白。", .95, 5.25, 10.2, .35, 17, TEAL, True)

# 3. first lesson error diagnosis
s = base("第 1 节｜易错点检测")
chip(s, "先找容易错的地方", .75, 1.0, 1.7, GOLD)
text(s, "请独立完成，不要着急说答案", .75, 1.65, 10.4, .45, 24, INK, True)
card(s, "E1 运算顺序", "18 − 3 × 4 = ____\n先算哪一步？", 1.0, 2.65, 4.7, 1.7, MINT)
card(s, "E2 数量关系", "3 盒饼干，每盒 4 块。\n一共有多少块？", 7.0, 2.65, 4.7, 1.7, ROSE, GOLD)
text(s, "写下答案，也写下你第一步做了什么。", 1.0, 5.25, 10.6, .35, 17, TEAL, True)

# 4. correction
s = base("第 1 节｜复习订正")
text(s, "先判断关系，再计算", .85, 1.15, 10.5, .45, 26, INK, True)
card(s, "第 1 题", "先算 3 × 4 = 12\n18 − 12 = 6", 1.0, 2.45, 4.7, 1.75, MINT)
card(s, "第 2 题", "每盒一样多：\n3 × 4 = 12（块）", 7.0, 2.45, 4.7, 1.75, ROSE, GOLD)
text(s, "检查：算式中的每个数，都要能在题目里找到意思。", 1.0, 5.2, 10.8, .35, 17, TEAL, True)

# 5. main example
s = base("第 1 节｜新内容经典例题")
text(s, "经典例题：两种买法，怎样求出单价？", .75, 1.1, 11.4, .5, 25, INK, True)
card(s, "购买 A｜27 元", "3 本练习本\n+ 2 支签字笔", 1.0, 2.2, 4.8, 2.1, MINT)
card(s, "购买 B｜35 元", "1 本练习本\n+ 5 支签字笔", 7.0, 2.2, 4.8, 2.1, ROSE, GOLD)
text(s, "问题：一本练习本和一支签字笔各多少元？", 1.0, 5.2, 10.6, .35, 18, TEAL, True)

# 6. silent thinking
s = base("静想一下")
text(s, "先不急着计算", 1.0, 1.3, 11.2, .5, 29, INK, True, PP_ALIGN.CENTER)
card(s, "试一试", "把两种购买信息写成两行。\n你想让哪一种物品的数量变得一样？", 1.3, 2.75, 4.6, 1.8, MINT)
card(s, "想一想", "如果两边都有同样多的练习本，\n相减后会留下什么？", 7.4, 2.75, 4.6, 1.8, ROSE, GOLD)
text(s, "把你的想法写在草稿纸上。", 1.0, 5.45, 11.2, .35, 18, TEAL, True, PP_ALIGN.CENTER)

# 7. analysis
s = base("第 1 节｜例题分析")
text(s, "先把练习本的数量对齐", .85, 1.1, 11.2, .48, 25, INK, True)
card(s, "购买 A", "3 本练习本\n+ 2 支笔 = 27 元", .95, 2.15, 3.35, 1.85, MINT)
card(s, "购买 B 的 3 倍", "3 本练习本\n+ 15 支笔 = 105 元", 4.95, 2.15, 3.35, 1.85, WHITE, GOLD)
card(s, "相减后", "练习本抵消，\n13 支笔 = 78 元", 8.95, 2.15, 3.35, 1.85, ROSE, GOLD)
text(s, "关键一步：先让一种物品的数量相同。", .95, 5.2, 11.1, .35, 18, TEAL, True)

# 8. solution
s = base("第 1 节｜规范解答与检验")
card(s, "签字笔", "105 − 27 = 78（元）\n78 ÷ 13 = 6（元）", 1.0, 1.9, 3.4, 2.2, MINT)
card(s, "练习本", "27 − 2 × 6 = 15（元）\n15 ÷ 3 = 5（元）", 4.95, 1.9, 3.4, 2.2, WHITE, GOLD)
card(s, "检验", "1 × 5 + 5 × 6\n= 35（元），符合购买 B", 8.9, 1.9, 3.4, 2.2, ROSE, GOLD)
text(s, "方法卡：先对齐一种数量，再相减，最后代回检验。", 1.0, 5.15, 10.8, .35, 18, TEAL, True)

# 9. lesson 2 test
s = base("第 2 节｜新内容检测")
text(s, "独立完成：这次不提示", .8, 1.15, 11.0, .45, 26, INK, True, PP_ALIGN.CENTER)
card(s, "购买 A", "2 本练习本 + 3 支笔\n一共 28 元", 1.0, 2.35, 4.8, 1.9, MINT)
card(s, "购买 B", "1 本练习本 + 5 支笔\n一共 35 元", 7.0, 2.35, 4.8, 1.9, ROSE, GOLD)
text(s, "问题：先把哪一条购买信息变成几倍？再求两种物品的单价。", 1.0, 5.15, 11.0, .35, 17, TEAL, True)

# 10. feedback
s = base("第 2 节｜检测反馈")
text(s, "先把购买 B 变成 2 倍", .8, 1.2, 11.0, .48, 26, INK, True, PP_ALIGN.CENTER)
card(s, "对齐", "2 × 35 = 70（元）\n2 本练习本 + 10 支笔", 1.0, 2.6, 4.7, 1.8, MINT)
card(s, "相减", "70 − 28 = 42（元）\n7 支笔 = 42 元", 4.35, 2.6, 4.7, 1.8, WHITE, GOLD)
card(s, "代回", "每支笔 6 元\n每本练习本 5 元", 7.7, 2.6, 4.7, 1.8, ROSE, GOLD)

# 11. transfer test
s = base("第 2 节｜迁移检测")
text(s, "不计算单价，先说“怎样对齐”", .8, 1.2, 11.0, .48, 25, INK, True, PP_ALIGN.CENTER)
text(s, "甲：4 张贴纸 + 2 支铅笔 = 32 元", 1.0, 2.35, 10.8, .35, 21, INK, True)
text(s, "乙：2 张贴纸 + 5 支铅笔 = 40 元", 1.0, 2.95, 10.8, .35, 21, INK, True)
card(s, "请回答", "应该把哪一条购买信息变成几倍？\n相减后会留下哪一种物品？", 2.0, 4.05, 9.3, 1.2, MINT)

# 12. transfer feedback
s = base("第 2 节｜迁移反馈")
text(s, "把乙变成 2 倍，让贴纸数量相同", 1.0, 1.35, 11.0, .5, 25, INK, True, PP_ALIGN.CENTER)
card(s, "乙的 2 倍", "4 张贴纸 + 10 支铅笔\n一共 80 元", 1.3, 2.8, 4.6, 1.8, MINT)
card(s, "相减后", "4 张贴纸抵消，\n留下 8 支铅笔的价钱", 7.4, 2.8, 4.6, 1.8, ROSE, GOLD)
text(s, "不管题目换成什么，先对齐、再相减的思路不变。", 1.0, 5.3, 11.0, .35, 17, TEAL, True, PP_ALIGN.CENTER)

# 13. preview
s = base("第 2 节｜课末预习")
chip(s, "只预测，不讲解", .7, 1.0, 1.45, GOLD)
text(s, "下一单元：两、三位数乘一位数", .7, 1.6, 10.8, .5, 27, INK, True)
text(s, "一盒鸡蛋有 20 个，买 3 盒一共有多少个？", 1.0, 2.45, 11.0, .4, 22, TEAL, True, PP_ALIGN.CENTER)
for i in range(3):
    shape(s, Shape.ROUNDED_RECTANGLE, 2.3+i*3.05, 3.35, 2.6, 1.5, MINT, MINT, True)
    text(s, "一盒\n20 个", 2.3+i*3.05, 3.85, 2.6, .6, 18, INK, True, PP_ALIGN.CENTER)
text(s, "预测问题：先猜一猜一共有多少个，再说说你打算怎样想 20 × 3；不用列竖式。", .9, 5.35, 11.3, .45, 16, TEAL, True, PP_ALIGN.CENTER)

# 14. close
s = base("本周方法")
text(s, "我会这样解决两种购买的问题：", .85, 1.3, 11.0, .5, 26, INK, True, PP_ALIGN.CENTER)
card(s, "1. 对齐", "让一种物品的数量相同", 1.0, 2.8, 3.3, 1.55, MINT)
card(s, "2. 相减", "找出另一种物品的总价", 5.0, 2.8, 3.3, 1.55, WHITE, GOLD)
card(s, "3. 检验", "把答案代回原题检查", 9.0, 2.8, 3.3, 1.55, ROSE, GOLD)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.core_properties.title = "第 1 周｜看见关系，发现规律"
prs.core_properties.subject = "三年级数学课堂演示稿"
prs.core_properties.author = "Codex"
prs.save(OUT)
print(f"created {OUT} with {len(prs.slides)} slides")
