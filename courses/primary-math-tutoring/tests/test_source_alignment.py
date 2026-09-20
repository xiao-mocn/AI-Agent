from pathlib import Path
import unittest

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
PPT = ROOT / "slides" / "0004-week-02-addition-subtraction-strategies.pptx"


def read(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def ppt_text(path: Path) -> str:
    presentation = Presentation(path)
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )


class SourceAlignmentTest(unittest.TestCase):
    def test_course_archive_keeps_school_review_short_and_thinking_topic_primary(self):
        syllabus = read(ROOT / "SYLLABUS.md")
        resources = read(ROOT / "RESOURCES.md")
        self.assertIn("课内只做短诊断和易错点复现", syllabus)
        self.assertIn("思维专题承担完整主例题、独立检测和迁移", syllabus)
        self.assertIn("数数图形 → 找找规律 → 加减巧算 → 巧添符号", resources)

    def test_current_deck_uses_addition_subtraction_as_the_only_main_topic(self):
        text = ppt_text(PPT)
        self.assertIn("加减巧算", text)
        self.assertIn("课内综合诊断", text)
        self.assertIn("数数图形｜复习热身", text)
        self.assertIn("巧添符号｜预热", text)
        self.assertIn("999 − 1 + 998 − 2", text)
        self.assertIn("第 2 节｜独立检测", text)
        self.assertIn("第 2 节｜课末预习", text)
        for forbidden in ("两、三位数乘一位数", "购买 A", "对齐一种数量", "教师提示", "商品价格消元"):
            self.assertNotIn(forbidden, text)

    def test_deck_has_one_main_example_and_answer_pages_follow_question_pages(self):
        presentation = Presentation(PPT)
        slides = [
            "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
            for slide in presentation.slides
        ]
        self.assertGreaterEqual(len(slides), 16)
        self.assertLessEqual(len(slides), 22)
        self.assertEqual(sum("主例题" in slide for slide in slides), 1)
        self.assertLess(
            next(i for i, slide in enumerate(slides) if "静想一下" in slide),
            next(i for i, slide in enumerate(slides) if "第 1 节｜看见结构" in slide),
        )
        self.assertLess(
            next(i for i, slide in enumerate(slides) if "第 2 节｜独立检测" in slide),
            next(i for i, slide in enumerate(slides) if "第 2 节｜检测反馈" in slide),
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
